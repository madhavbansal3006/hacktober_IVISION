import csv

import openpyxl
import pytest

from src.schema import SCHEMA_VERSION, GeoCodedResult
from src.teammate_metadata_writer import (
    PHOTO_MASTER_FIELDNAMES,
    QA_LOG_FIELDNAMES,
    write_photo_master_validated_csv,
    write_processing_log,
    write_qa_log_xlsx,
)


def _make_result(**overrides) -> GeoCodedResult:
    defaults = dict(
        schema_version=SCHEMA_VERSION,
        image_name="photo_001.jpg",
        processing_timestamp="2026-08-31T10:15:00+00:00",
        latitude=15.01234,
        longitude=75.21345,
        status="OK",
        error_reason=None,
        inside_study_area=True,
        micro_watershed_id="g_mws.240575",
        mws_code="4D4F2m3",
        lulc_value=2,
        lulc_class="Agriculture",
        raster_pixel_row=1842,
        raster_pixel_col=2210,
        lulc_neighborhood_window_size=3,
        lulc_neighborhood_class_purity_pct=88.89,
        lulc_neighborhood_nodata_pct=0.0,
        image_id="IMG_0001",
        file_path="/data/photos/photo_001.jpg",
        capture_datetime="2026-03-14T09:21:07",
        gps_source="EXIF",
        gps_accuracy_m=4.5,
        coordinate_status="VALID",
        study_area_status="INSIDE",
    )
    defaults.update(overrides)
    return GeoCodedResult(**defaults)


def _make_failure_result(**overrides) -> GeoCodedResult:
    defaults = dict(
        schema_version=SCHEMA_VERSION,
        image_name="no_gps.jpg",
        processing_timestamp="2026-08-31T10:16:00+00:00",
        latitude=None,
        longitude=None,
        status="NO_GPS",
        error_reason="No GPS IFD present in EXIF metadata",
        inside_study_area=None,
        micro_watershed_id=None,
        mws_code=None,
        lulc_value=None,
        lulc_class=None,
        raster_pixel_row=None,
        raster_pixel_col=None,
        lulc_neighborhood_window_size=None,
        lulc_neighborhood_class_purity_pct=None,
        lulc_neighborhood_nodata_pct=None,
        image_id="IMG_0002",
        file_path="/data/photos/no_gps.jpg",
        capture_datetime=None,
        gps_source=None,
        gps_accuracy_m=None,
        coordinate_status="MISSING",
        study_area_status=None,
    )
    defaults.update(overrides)
    return GeoCodedResult(**defaults)


class TestPhotoMasterValidatedCsv:
    def test_writes_expected_header(self, tmp_path):
        path = write_photo_master_validated_csv([_make_result()], tmp_path / "pmv.csv")
        with open(path, newline="") as f:
            header = next(csv.reader(f))
        assert header == PHOTO_MASTER_FIELDNAMES

    def test_row_count_matches_input(self, tmp_path):
        results = [_make_result(), _make_failure_result()]
        path = write_photo_master_validated_csv(results, tmp_path / "pmv.csv")
        with open(path, newline="") as f:
            rows = list(csv.DictReader(f))
        assert len(rows) == 2

    def test_mws_id_column_populated_from_micro_watershed_id(self, tmp_path):
        """Explicit integration-key check: mws_id in the CSV must equal
        micro_watershed_id (the `id` field), not mws_code (`MWS`)."""
        result = _make_result(micro_watershed_id="g_mws.999999", mws_code="XYZ123")
        path = write_photo_master_validated_csv([result], tmp_path / "pmv.csv")
        with open(path, newline="") as f:
            row = next(csv.DictReader(f))
        assert row["mws_id"] == "g_mws.999999"
        assert row["mws_id"] != "XYZ123"

    def test_human_review_fields_always_blank(self, tmp_path):
        result = _make_result()
        path = write_photo_master_validated_csv([result], tmp_path / "pmv.csv")
        with open(path, newline="") as f:
            row = next(csv.DictReader(f))
        for field in (
            "observation_type", "description", "photo_quality",
            "evidence_relevance", "intervention_type", "remarks",
            "review_status", "reviewer",
        ):
            assert row[field] == "", f"{field} must be blank, got {row[field]!r}"

    def test_failure_result_row_has_correct_fields(self, tmp_path):
        result = _make_failure_result()
        path = write_photo_master_validated_csv([result], tmp_path / "pmv.csv")
        with open(path, newline="") as f:
            row = next(csv.DictReader(f))
        assert row["image_id"] == "IMG_0002"
        assert row["coordinate_status"] == "MISSING"
        assert row["latitude"] == ""
        assert row["mws_id"] == ""

    def test_creates_parent_directories(self, tmp_path):
        nested = tmp_path / "a" / "b" / "Photo_Master_Validated.csv"
        write_photo_master_validated_csv([_make_result()], nested)
        assert nested.exists()

    def test_all_expected_field_values_correct(self, tmp_path):
        result = _make_result()
        path = write_photo_master_validated_csv([result], tmp_path / "pmv.csv")
        with open(path, newline="") as f:
            row = next(csv.DictReader(f))
        assert row["image_id"] == "IMG_0001"
        assert row["file_name"] == "photo_001.jpg"
        assert row["file_path"] == "/data/photos/photo_001.jpg"
        assert float(row["latitude"]) == pytest.approx(15.01234)
        assert row["capture_datetime"] == "2026-03-14T09:21:07"
        assert row["gps_source"] == "EXIF"
        assert float(row["gps_accuracy_m"]) == pytest.approx(4.5)
        assert row["coordinate_status"] == "VALID"
        assert row["study_area_status"] == "INSIDE"


class TestQaLogXlsx:
    def test_writes_workbook_with_correct_header(self, tmp_path):
        path = write_qa_log_xlsx([_make_result()], tmp_path / "qa.xlsx")
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        header = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        assert header == QA_LOG_FIELDNAMES

    def test_row_count_matches_input_plus_header(self, tmp_path):
        results = [_make_result(), _make_failure_result()]
        path = write_qa_log_xlsx(results, tmp_path / "qa.xlsx")
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        assert ws.max_row == 3  # header + 2 data rows

    def test_row_values_correct(self, tmp_path):
        result = _make_result()
        path = write_qa_log_xlsx([result], tmp_path / "qa.xlsx")
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        data_row = [cell.value for cell in next(ws.iter_rows(min_row=2, max_row=2))]
        row_dict = dict(zip(QA_LOG_FIELDNAMES, data_row))
        assert row_dict["image_id"] == "IMG_0001"
        assert row_dict["mws_id"] == "g_mws.240575"
        assert row_dict["lulc_class"] == "Agriculture"
        assert row_dict["status"] == "OK"
        assert row_dict["review_status"] is None
        assert row_dict["remarks"] is None

    def test_creates_parent_directories(self, tmp_path):
        nested = tmp_path / "qa" / "sub" / "QA_Log.xlsx"
        write_qa_log_xlsx([_make_result()], nested)
        assert nested.exists()

    def test_empty_results_still_writes_header_only(self, tmp_path):
        path = write_qa_log_xlsx([], tmp_path / "qa.xlsx")
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        assert ws.max_row == 1


class TestProcessingLog:
    def test_counts_are_correct(self, tmp_path):
        results = [
            _make_result(),  # OK, geotagged, inside, watershed matched, lulc matched
            _make_failure_result(),  # NO_GPS -> MISSING
            _make_result(
                image_id="IMG_0003", coordinate_status="SUSPICIOUS",
                study_area_status=None, status="INVALID_COORDINATE",
                latitude=None, longitude=None, micro_watershed_id=None,
                mws_code=None, lulc_value=None, lulc_class=None,
                error_reason="bad coordinate",
            ),
            _make_result(
                image_id="IMG_0004", status="OUTSIDE_STUDY_AREA",
                coordinate_status="OUTSIDE_STUDY_AREA", study_area_status="OUTSIDE",
                inside_study_area=False, micro_watershed_id=None, mws_code=None,
                lulc_value=None, lulc_class=None,
            ),
            _make_result(
                image_id="IMG_0005", status="OK",
                study_area_status="BOUNDARY_REVIEW",
            ),
        ]
        path = write_processing_log(results, tmp_path / "log.txt")
        content = path.read_text()

        assert "TOTAL_IMAGES: 5" in content
        assert "GEOTAGGED_IMAGES: 3" in content  # results with lat/lon set
        assert "MISSING_GPS: 1" in content
        assert "INVALID_MALFORMED_GPS: 1" in content
        assert "OUTSIDE_STUDY_AREA: 1" in content
        assert "BOUNDARY_REVIEW: 1" in content
        assert "ERRORS: 2" in content  # no_gps + invalid_coordinate both have error_reason

    def test_includes_crs_and_mws_field_info(self, tmp_path):
        path = write_processing_log([_make_result()], tmp_path / "log.txt")
        content = path.read_text()
        assert "CRS_USED" in content
        assert "MWS_FIELD_USED: id" in content
        assert "BOUNDARY_REVIEW_TOLERANCE_M" in content

    def test_includes_processing_timestamp(self, tmp_path):
        path = write_processing_log([_make_result()], tmp_path / "log.txt")
        content = path.read_text()
        assert "PROCESSING_TIMESTAMP:" in content

    def test_empty_results_produces_zero_counts(self, tmp_path):
        path = write_processing_log([], tmp_path / "log.txt")
        content = path.read_text()
        assert "TOTAL_IMAGES: 0" in content
        assert "GEOTAGGED_IMAGES: 0" in content

    def test_creates_parent_directories(self, tmp_path):
        nested = tmp_path / "logs" / "sub" / "log.txt"
        write_processing_log([_make_result()], nested)
        assert nested.exists()
